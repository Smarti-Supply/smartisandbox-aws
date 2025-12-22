import base64
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def generate_excel(data: list[dict]) -> bytes:
    """
    Gera arquivo Excel a partir de array de objetos consolidados.
    
    Args:
        data: Lista de dicionários com os dados consolidados
    
    Returns:
        bytes: Arquivo Excel em bytes
    """
    if not data or len(data) == 0:
        raise ValueError("Dados não podem estar vazios")
    
    # Criar workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Order Details"
    
    # Definir ordem das colunas (campos específicos primeiro, depois os demais)
    preferred_order = [
        'id_pedido',
        'numero_item',
        # Campos de order
        'numero_pedido',
        'pedido_descricao',
        'pedido_data_da_remessa',
        'fornecedor',
        'pedido_criado_em',
        'pedido_atualizado_em',
        # Campos de order_item
        'numero_item',
        'produto',
        'descricao_produto',
        'quantidade',
        'unidade_medida',
        'preco_unitario',
        'preco_total',
        'centro',
        'item_data_da_remessa',
        'data_da_entrega',
        'status',
        'item_criado_em',
        # Campos consolidados de observations e followup_tracking (com nomes em português)
        'observacao do usuario',
        'observacao do fornecedor',
        'regras de followup',
        # Campos consolidados de order_item_invoices
        'numero_nfe',
        'data_nfe',
        'quantidade_faturada',
        'quantidade_faturada_criado_em'
    ]
    
    # Extrair colunas: primeiro as preferidas na ordem, depois as restantes ordenadas
    all_keys = set()
    for row in data:
        all_keys.update(row.keys())
    
    # Ordenar colunas: preferidas primeiro, depois as restantes
    columns = []
    for col in preferred_order:
        if col in all_keys:
            columns.append(col)
            all_keys.remove(col)
    
    # Adicionar as restantes ordenadas
    columns.extend(sorted(all_keys))
    
    # Escrever header (primeira linha)
    header_font = Font(bold=True)
    for col_idx, column in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=column)
        cell.font = header_font
    
    # Preencher dados linha por linha
    for row_idx, row_data in enumerate(data, 2):  # Começar na linha 2
        for col_idx, column in enumerate(columns, 1):
            value = row_data.get(column)
            
            # Tratar valores None como string vazia
            if value is None:
                value = ""
            
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    # Ajustar largura das colunas automaticamente
    for col_idx, column in enumerate(columns, 1):
        column_letter = get_column_letter(col_idx)
        max_length = 0
        
        # Verificar header
        max_length = max(max_length, len(str(column)))
        
        # Verificar dados
        for row_data in data:
            value = row_data.get(column)
            if value is not None:
                max_length = max(max_length, len(str(value)))
        
        # Definir largura (mínimo 10, máximo 50)
        adjusted_width = min(max(max_length + 2, 10), 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Converter para bytes
    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    
    return excel_buffer.getvalue()

