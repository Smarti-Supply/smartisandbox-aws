import base64
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def generate_excel(data: list[dict], table_name: str) -> bytes:
    """
    Gera arquivo Excel genérico a partir de array de objetos.
    
    Args:
        data: Lista de dicionários com os dados
        table_name: Nome da tabela (usado para nome da sheet)
    
    Returns:
        bytes: Arquivo Excel em bytes
    """
    if not data or len(data) == 0:
        raise ValueError("Dados não podem estar vazios")
    
    # Criar workbook
    wb = Workbook()
    ws = wb.active
    ws.title = table_name.capitalize()
    
    # Extrair colunas das chaves do primeiro objeto
    columns = sorted(data[0].keys())
    
    # Escrever header (primeira linha)
    header_font = Font(bold=True)
    for col_idx, column in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=column)
        cell.font = header_font
    
    # Preencher dados linha por linha
    for row_idx, row_data in enumerate(data, 2):  # Começar na linha 2
        for col_idx, column in enumerate(columns, 1):
            value = row_data.get(column)
            
            # Tratar valores None como string vazia
            if value is None:
                value = ""
            
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    # Ajustar largura das colunas automaticamente
    for col_idx, column in enumerate(columns, 1):
        column_letter = get_column_letter(col_idx)
        max_length = 0
        
        # Verificar header
        max_length = max(max_length, len(str(column)))
        
        # Verificar dados
        for row_data in data:
            value = row_data.get(column)
            if value is not None:
                max_length = max(max_length, len(str(value)))
        
        # Definir largura (mínimo 10, máximo 50)
        adjusted_width = min(max(max_length + 2, 10), 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Converter para bytes
    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    
    return excel_buffer.getvalue()
